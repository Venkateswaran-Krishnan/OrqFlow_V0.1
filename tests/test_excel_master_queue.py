from __future__ import annotations

import sqlite3
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from framework.runtime.queue_db import SQLiteQueueDatabase, load_queue_queries
from framework.runtime.queue_runtime import create_master_queue, get_next_transaction


REPO_ROOT = Path(__file__).resolve().parents[1]
SHARE_ROOT = REPO_ROOT / "share"


class SharedQueueQueryTests(unittest.TestCase):
    def test_sqlite_query_file_loads(self) -> None:
        queries = load_queue_queries("sqlite", {"share_root": str(SHARE_ROOT)})

        self.assertIn("tbl_input_columns", queries)
        self.assertIn("insert_tbl_input", queries)
        self.assertIn("insert_tbl_queue", queries)
        self.assertIn("select_inputs_for_queue_creation", queries)
        self.assertIn("mark_input_queue_created", queries)
        self.assertEqual("PRAGMA table_info(tbl_input);", queries["tbl_input_columns"])
        self.assertEqual(
            "INSERT INTO tbl_input ({columns})\nVALUES ({placeholders});",
            queries["insert_tbl_input"],
        )

    def test_mysql_query_file_loads_without_connection(self) -> None:
        queries = load_queue_queries("mysql", {"share_root": str(SHARE_ROOT)})

        self.assertIn("tbl_input_columns", queries)
        self.assertIn("insert_tbl_input", queries)
        self.assertIn("insert_tbl_queue", queries)
        self.assertIn("select_inputs_for_queue_creation", queries)
        self.assertIn("mark_input_queue_created", queries)
        self.assertEqual("SHOW COLUMNS FROM tbl_input;", queries["tbl_input_columns"])
        self.assertEqual(
            "INSERT INTO tbl_input ({columns})\nVALUES ({placeholders});",
            queries["insert_tbl_input"],
        )

    def test_sqlite_and_mysql_query_files_have_same_names(self) -> None:
        sqlite_queries = load_queue_queries("sqlite", {"share_root": str(SHARE_ROOT)})
        mysql_queries = load_queue_queries("mysql", {"share_root": str(SHARE_ROOT)})

        self.assertEqual(set(sqlite_queries), set(mysql_queries))


class ExcelMasterQueueTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.project_dir = Path(self.temp_dir.name)
        self.db_path = self.project_dir / "queue.sqlite3"
        self._create_schema()
        self.db = SQLiteQueueDatabase(
            path=self.db_path,
            config={"type": "sqlite", "sqlite_path": str(self.db_path)},
            queries=load_queue_queries("sqlite", {"share_root": str(SHARE_ROOT)}),
        )
        self.db.connect()

    def tearDown(self) -> None:
        self.db.close()
        self.temp_dir.cleanup()

    def test_masterbot_false_skips_excel_loading(self) -> None:
        state = self._state(masterbot=False)

        result = create_master_queue(state)

        self.assertIs(result, state)
        self.assertEqual(0, self._count("tbl_input"))

    def test_api_source_uses_api_adapter_placeholder(self) -> None:
        state = self._state(queue_source="API")

        create_master_queue(state)

        self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
        self.assertIn("API input loading is not configured yet", state["runtime_config"]["last_error"])
        self.assertEqual(0, self._count("tbl_input"))

    def test_missing_mandatory_column_fails_before_insert(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Process", "Chargeback_Date", "Case_Json"],
            [[1, 2, "2026-05-25", "{}"]],
        )
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)

        self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
        self.assertIn("Case_ID", state["runtime_config"]["last_error"])
        self.assertEqual("END", state["runtime_config"]["next_action"])
        self.assertEqual(0, self._count("tbl_input"))

    def test_valid_excel_rows_insert_into_input_and_distinct_application_queues(self) -> None:
        workbook_path = self._write_excel(
            [
                "Processor",
                "Chargeback_Date",
                "Case_Json",
                "Case_ID",
                "Ignored_Column",
            ],
            [
                [10, "2026-05-25", '{"case": "A"}', "CASE-1", "ignore me"],
            ],
        )
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)

        rows = self.db.connection.execute(
            """
            SELECT Processor, Process, Case_ID, Input_Identifier, Status,
                   QueueCreation_timestamp
            FROM tbl_input
            """
        ).fetchall()
        self.assertEqual(1, len(rows))
        self.assertEqual(10, rows[0]["Processor"])
        self.assertEqual(12, rows[0]["Process"])
        self.assertEqual("CASE-1", rows[0]["Case_ID"])
        self.assertEqual("12_CASE-1", rows[0]["Input_Identifier"])
        self.assertEqual("Queue Created", rows[0]["Status"])
        self.assertIsNotNone(rows[0]["QueueCreation_timestamp"])
        queue_rows = self.db.connection.execute(
            """
            SELECT Case_Details, Application_Details, Processing_Status, Bot_Name
            FROM tbl_queue
            ORDER BY ID
            """
        ).fetchall()
        self.assertEqual(2, len(queue_rows))
        self.assertEqual([13, 14], [row["Application_Details"] for row in queue_rows])
        self.assertTrue(all(row["Case_Details"] == 1 for row in queue_rows))
        self.assertTrue(all(row["Processing_Status"] == "Queue Created" for row in queue_rows))
        self.assertTrue(all(row["Bot_Name"] is None for row in queue_rows))
        self.assertEqual(1, state["runtime_config"]["input_load_summary"]["inserted_count"])
        self.assertEqual(
            {
                "distinct_application_count": 2,
                "eligible_input_count": 1,
                "queued_input_count": 1,
                "created_queue_count": 2,
                "failed_input_count": 0,
                "failed_inputs": [],
            },
            state["runtime_config"]["queue_creation_summary"],
        )

    def test_source_process_column_is_ignored(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Process", "Chargeback_Date", "Case_Json", "Case_ID"],
            [[10, 999, "2026-05-25", "{}", "CASE-1"]],
        )
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)

        row = self.db.connection.execute(
            "SELECT Process, Input_Identifier FROM tbl_input"
        ).fetchone()
        self.assertEqual(12, row["Process"])
        self.assertEqual("12_CASE-1", row["Input_Identifier"])

    def test_missing_process_id_fails_before_insert(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [[10, "2026-05-25", "CASE-1"]],
        )
        state = self._state(queue_file=workbook_path.name, process_id="")

        create_master_queue(state)

        self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
        self.assertIn("process_config.Process_ID", state["runtime_config"]["last_error"])
        self.assertEqual(0, self._count("tbl_input"))

    def test_missing_required_value_skips_only_that_row(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [
                [10, "2026-05-25", "CASE-1"],
                [None, "2026-05-25", "CASE-2"],
                [10, "2026-05-25", "CASE-3"],
            ],
        )
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)

        summary = state["runtime_config"]["input_load_summary"]
        self.assertEqual(2, self._count("tbl_input"))
        self.assertEqual(2, summary["inserted_count"])
        self.assertEqual(1, summary["failed_count"])
        self.assertEqual(3, summary["failed_rows"][0]["row"])
        self.assertIn("Processor", summary["failed_rows"][0]["error"])

    def test_duplicate_input_identifier_skips_row_and_continues(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [
                [10, "2026-05-25", "CASE-1"],
                [10, "2026-05-25", "CASE-1"],
                [10, "2026-05-25", "CASE-2"],
            ],
        )
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)

        summary = state["runtime_config"]["input_load_summary"]
        rows = self.db.connection.execute(
            "SELECT Case_ID, Input_Identifier, Status FROM tbl_input ORDER BY ID"
        ).fetchall()
        self.assertEqual(2, len(rows))
        self.assertEqual(["CASE-1", "CASE-2"], [row["Case_ID"] for row in rows])
        self.assertEqual(["12_CASE-1", "12_CASE-2"], [row["Input_Identifier"] for row in rows])
        self.assertTrue(all(row["Status"] == "Queue Created" for row in rows))
        self.assertEqual(2, summary["inserted_count"])
        self.assertEqual(1, summary["failed_count"])
        self.assertIn("DB insert failed", summary["failed_rows"][0]["error"])
        self.assertEqual(4, self._count("tbl_queue"))

    def test_distinct_applications_follow_first_sequence_order(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [[10, "2026-05-25", "CASE-1"]],
        )
        state = self._state(queue_file=workbook_path.name)
        state["key_steps"] = self._key_steps(
            [(30, 13), (10, "14"), (20, 13.0), (40, 15)]
        )

        create_master_queue(state)

        rows = self.db.connection.execute(
            "SELECT Application_Details FROM tbl_queue ORDER BY ID"
        ).fetchall()
        self.assertEqual([14, 13, 15], [row["Application_Details"] for row in rows])
        self.assertEqual(
            3,
            state["runtime_config"]["queue_creation_summary"]["distinct_application_count"],
        )

    def test_only_matching_process_and_blank_status_inputs_are_queued(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [],
        )
        eligible_ids = [
            self._insert_input(process=12, case_id="ELIGIBLE-NULL", status=None),
            self._insert_input(process=12, case_id="ELIGIBLE-EMPTY", status=""),
            self._insert_input(process=12, case_id="ELIGIBLE-SPACE", status="   "),
        ]
        other_process_id = self._insert_input(process=99, case_id="OTHER", status=None)
        completed_id = self._insert_input(
            process=12,
            case_id="DONE",
            status="Already Queued",
        )
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)

        rows = self.db.connection.execute(
            "SELECT Case_Details, Application_Details FROM tbl_queue ORDER BY ID"
        ).fetchall()
        self.assertEqual(6, len(rows))
        self.assertEqual(
            [eligible_ids[0], eligible_ids[0], eligible_ids[1], eligible_ids[1], eligible_ids[2], eligible_ids[2]],
            [row["Case_Details"] for row in rows],
        )
        untouched = self.db.connection.execute(
            "SELECT ID, Status FROM tbl_input WHERE ID IN (?, ?) ORDER BY ID",
            [other_process_id, completed_id],
        ).fetchall()
        self.assertIsNone(untouched[0]["Status"])
        self.assertEqual("Already Queued", untouched[1]["Status"])

    def test_rerun_does_not_duplicate_queues(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [[10, "2026-05-25", "CASE-1"]],
        )
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)
        create_master_queue(state)

        self.assertEqual(2, self._count("tbl_queue"))
        self.assertEqual(
            0,
            state["runtime_config"]["queue_creation_summary"]["eligible_input_count"],
        )

    def test_configured_eligible_status_is_used_for_input_and_queue(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [[10, "2026-05-25", "CASE-1"]],
        )
        state = self._state(queue_file=workbook_path.name)
        state["config"]["queue_config"]["eligible_status"] = "Ready for Bot"

        create_master_queue(state)

        input_status = self.db.connection.execute(
            "SELECT Status FROM tbl_input"
        ).fetchone()["Status"]
        queue_statuses = self.db.connection.execute(
            "SELECT Processing_Status FROM tbl_queue ORDER BY ID"
        ).fetchall()
        self.assertEqual("Ready for Bot", input_status)
        self.assertEqual(
            ["Ready for Bot", "Ready for Bot"],
            [row["Processing_Status"] for row in queue_statuses],
        )

    def test_invalid_key_step_application_fails_before_input_writes(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [[10, "2026-05-25", "CASE-1"]],
        )
        state = self._state(queue_file=workbook_path.name)
        state["key_steps"] = self._key_steps([(1, 999)])

        create_master_queue(state)

        self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
        self.assertIn("999", state["runtime_config"]["last_error"])
        self.assertEqual(0, self._count("tbl_input"))
        self.assertEqual(0, self._count("tbl_queue"))

    def test_blank_nonnumeric_and_nonpositive_key_step_applications_fail(self) -> None:
        for application in (None, "not-an-id", 0, -1, 13.5):
            with self.subTest(application=application):
                workbook_path = self._write_excel(
                    ["Processor", "Chargeback_Date", "Case_ID"],
                    [[10, "2026-05-25", "CASE-1"]],
                )
                state = self._state(queue_file=workbook_path.name)
                state["key_steps"] = self._key_steps([(1, application)])

                create_master_queue(state)

                self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
                self.assertIn("Application", state["runtime_config"]["last_error"])
                self.assertEqual(0, self._count("tbl_input"))
                self.assertEqual(0, self._count("tbl_queue"))

    def test_missing_process_key_steps_fails_before_input_writes(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [[10, "2026-05-25", "CASE-1"]],
        )
        state = self._state(queue_file=workbook_path.name)
        state["key_steps"] = pd.DataFrame(
            [{"Sequence": 1, "State": "FRAMEWORK_INIT", "Application": 13}]
        )

        create_master_queue(state)

        self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
        self.assertIn("no PROCESS_TRANSACTION", state["runtime_config"]["last_error"])
        self.assertEqual(0, self._count("tbl_input"))

    def test_queue_insert_failure_rolls_back_input_set_and_continues(self) -> None:
        workbook_path = self._write_excel(
            ["Processor", "Chargeback_Date", "Case_ID"],
            [
                [10, "2026-05-25", "CASE-1"],
                [10, "2026-05-25", "CASE-2"],
            ],
        )
        self.db.connection.execute(
            """
            CREATE TRIGGER fail_second_application_for_first_input
            BEFORE INSERT ON tbl_queue
            WHEN NEW.Case_Details = 1 AND NEW.Application_Details = 14
            BEGIN
                SELECT RAISE(ABORT, 'simulated queue failure');
            END
            """
        )
        self.db.connection.commit()
        state = self._state(queue_file=workbook_path.name)

        create_master_queue(state)

        summary = state["runtime_config"]["queue_creation_summary"]
        first_input = self.db.connection.execute(
            "SELECT Status, QueueCreation_timestamp FROM tbl_input WHERE ID = 1"
        ).fetchone()
        second_input = self.db.connection.execute(
            "SELECT Status, QueueCreation_timestamp FROM tbl_input WHERE ID = 2"
        ).fetchone()
        queue_rows = self.db.connection.execute(
            "SELECT Case_Details, Application_Details FROM tbl_queue ORDER BY ID"
        ).fetchall()
        self.assertIsNone(first_input["Status"])
        self.assertIsNone(first_input["QueueCreation_timestamp"])
        self.assertEqual("Queue Created", second_input["Status"])
        self.assertIsNotNone(second_input["QueueCreation_timestamp"])
        self.assertEqual([(2, 13), (2, 14)], [tuple(row) for row in queue_rows])
        self.assertEqual(1, summary["queued_input_count"])
        self.assertEqual(2, summary["created_queue_count"])
        self.assertEqual(1, summary["failed_input_count"])
        self.assertEqual(1, summary["failed_inputs"][0]["input_id"])
        self.assertIn("simulated queue failure", summary["failed_inputs"][0]["error"])

    def test_tbl_input_columns_query_returns_usable_sqlite_columns(self) -> None:
        rows = self.db.connection.execute(self.db.queries["tbl_input_columns"]).fetchall()
        names = {row["name"] for row in rows}

        self.assertIn("Case_ID", names)
        self.assertIn("Input_Identifier", names)

    def test_get_next_transaction_uses_sqlite_queue_and_marks_in_progress(self) -> None:
        self._insert_input_and_queue("Queue Created")
        state = self._state()

        get_next_transaction(state)

        txn = state["runtime_config"]["txn"]
        self.assertEqual("PROCESS", state["runtime_config"]["next_action"])
        self.assertEqual("In Processing", txn["queue_processing_status"])
        self.assertEqual({"case": "CASE-1"}, txn["case_json"])
        queue_row = self.db.connection.execute(
            "SELECT Processing_Status, Bot_Name, ProcessingSTART_timestamp FROM tbl_queue"
        ).fetchone()
        self.assertEqual("In Processing", queue_row["Processing_Status"])
        self.assertEqual("BOT-1", queue_row["Bot_Name"])
        self.assertIsNotNone(queue_row["ProcessingSTART_timestamp"])

    def test_get_next_transaction_returns_no_transaction_when_no_eligible_queue_row(self) -> None:
        self._insert_input_and_queue("Already Done")
        state = self._state()

        get_next_transaction(state)

        self.assertIsNone(state["runtime_config"]["txn"])
        self.assertEqual("NO_TRANSACTION", state["runtime_config"]["last_status"])

    def test_get_next_transaction_missing_queue_db_fails_gracefully(self) -> None:
        state = self._state()
        state.pop("queue_db")

        get_next_transaction(state)

        self.assertIsNone(state["runtime_config"]["txn"])
        self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
        self.assertIn("queue_db", state["runtime_config"]["last_error"])
        self.assertIsNone(state["runtime_config"]["next_action"])

    def test_get_next_transaction_query_error_fails_gracefully(self) -> None:
        state = self._state()
        self.db.queries["fetch_next_transaction"] = "SELECT * FROM missing_table"

        get_next_transaction(state)

        self.assertIsNone(state["runtime_config"]["txn"])
        self.assertEqual("SYSTEM_EXCEPTION", state["runtime_config"]["last_status"])
        self.assertIn("missing_table", state["runtime_config"]["last_error"])
        self.assertIsNone(state["runtime_config"]["next_action"])

    def test_database_queue_writes_success_status(self) -> None:
        self._insert_input_and_queue("Queue Created")
        state = self._state()
        get_next_transaction(state)

        state["queue"].mark_success(state["runtime_config"]["txn"])

        queue_row = self.db.connection.execute(
            "SELECT Processing_Status, Bot_Comment, ProcessingEND_timestamp FROM tbl_queue"
        ).fetchone()
        self.assertEqual("Success", queue_row["Processing_Status"])
        self.assertIsNone(queue_row["Bot_Comment"])
        self.assertIsNotNone(queue_row["ProcessingEND_timestamp"])

    def test_success_with_no_reason_preserves_existing_bot_comment(self) -> None:
        self._insert_input_and_queue("Queue Created", bot_comment="existing note")
        state = self._state()
        get_next_transaction(state)

        state["queue"].mark_success(state["runtime_config"]["txn"])

        queue_row = self.db.connection.execute(
            "SELECT Processing_Status, Bot_Comment FROM tbl_queue"
        ).fetchone()
        self.assertEqual("Success", queue_row["Processing_Status"])
        self.assertEqual("existing note", queue_row["Bot_Comment"])

    def test_database_queue_writes_failed_and_skipped_statuses(self) -> None:
        input_id = self._insert_input_and_queue("Queue Created", bot_comment="first note")
        state = self._state()
        get_next_transaction(state)
        state["queue"].mark_failed(state["runtime_config"]["txn"], "failed reason")

        self.db.connection.execute(
            """
            INSERT INTO tbl_queue (Case_Details, Application_Details, Processing_Status)
            VALUES (?, ?, ?)
            """,
            [input_id, 10, "Queue Created"],
        )
        self.db.connection.commit()
        state["runtime_config"]["queue_initialized"] = False
        get_next_transaction(state)
        state["queue"].mark_skipped(state["runtime_config"]["txn"], "skip reason")

        rows = self.db.connection.execute(
            "SELECT Processing_Status, Bot_Comment FROM tbl_queue ORDER BY ID"
        ).fetchall()
        self.assertEqual("Failed", rows[0]["Processing_Status"])
        self.assertEqual("first note\nfailed reason", rows[0]["Bot_Comment"])
        self.assertEqual("Skipped", rows[1]["Processing_Status"])
        self.assertEqual("skip reason", rows[1]["Bot_Comment"])

    def _create_schema(self) -> None:
        connection = sqlite3.connect(self.db_path)
        try:
            connection.execute(
                """
                CREATE TABLE tbl_application (
                    ID INTEGER PRIMARY KEY,
                    App_Name TEXT
                )
                """
            )
            connection.executemany(
                "INSERT INTO tbl_application (ID, App_Name) VALUES (?, ?)",
                [(10, "Legacy"), (13, "Chase"), (14, "Wells Fargo"), (15, "Amex")],
            )
            connection.execute(
                """
                CREATE TABLE tbl_input (
                    ID INTEGER PRIMARY KEY AUTOINCREMENT,
                    Processor INTEGER NOT NULL,
                    Process INTEGER NOT NULL,
                    Case_Number TEXT,
                    Transaction_ID TEXT,
                    Case_Status TEXT,
                    Chargeback_Date TEXT NOT NULL,
                    Case_Json TEXT,
                    Transaction_Amount TEXT,
                    Mid_Alias TEXT,
                    MID_Number TEXT,
                    Case_ID TEXT NOT NULL,
                    Chargeback_Amount TEXT,
                    Transaction_Date TEXT,
                    Deadline_Date TEXT,
                    Card_First_Six TEXT,
                    Card_Last_Four TEXT,
                    Card_Type TEXT,
                    Status TEXT,
                    QueueCreation_timestamp TEXT,
                    Input_Identifier TEXT UNIQUE,
                    Institution TEXT,
                    BUnit TEXT,
                    Optional_Match TEXT
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE tbl_queue (
                    ID INTEGER PRIMARY KEY AUTOINCREMENT,
                    Case_Details INTEGER NOT NULL,
                    Application_Details INTEGER NOT NULL,
                    Bot_Name TEXT,
                    Processing_Status TEXT,
                    CTO_Details TEXT,
                    Evidence_Status TEXT,
                    Output_tbl_Status INTEGER,
                    Bot_Comment TEXT,
                    Dependency TEXT,
                    ProcessingSTART_timestamp TEXT,
                    ProcessingEND_timestamp TEXT
                )
                """
            )
            connection.commit()
        finally:
            connection.close()

    def _write_excel(self, headers: list[str], rows: list[list[object]]) -> Path:
        path = self.project_dir / "input.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def _state(
        self,
        masterbot: bool = True,
        queue_file: str = "input.xlsx",
        process_id: str = "12",
        queue_source: str = "Excel",
    ) -> dict:
        return {
            "config": {
                "queue_config": {
                    "eligible_status": "Queue Created",
                    "in_progress_status": "In Processing",
                    "success_status": "Success",
                    "failed_status": "Failed",
                    "skipped_status": "Skipped",
                },
                "process_config": {
                    "Process_ID": process_id,
                    "settings": {
                        "masterbot": masterbot,
                        "Queue": queue_source,
                        "QueueFileLocation": queue_file,
                        "ApiConfig": {},
                        "bot": "BOT-1",
                    }
                }
            },
            "config_context": {
                "share_root": str(SHARE_ROOT),
                "project_config_dir": str(self.project_dir),
            },
            "runtime_config": {},
            "queue_db": self.db,
            "key_steps": self._key_steps([(1, 13), (2, 14), (3, 13)]),
            "logs": [],
        }

    def _key_steps(self, sequence_applications: list[tuple[object, object]]) -> pd.DataFrame:
        return pd.DataFrame(
            [
                {
                    "Sequence": sequence,
                    "Bot": None,
                    "State": "PROCESS_TRANSACTION",
                    "BatchCount": None,
                    "Application": application,
                    "Moduel": None,
                    "Status": None,
                }
                for sequence, application in sequence_applications
            ]
        )

    def _count(self, table: str) -> int:
        row = self.db.connection.execute(f"SELECT COUNT(*) AS count FROM {table}").fetchone()
        return int(row["count"])

    def _insert_input_and_queue(self, status: str, bot_comment: str | None = None) -> int:
        cursor = self.db.connection.execute(
            """
            INSERT INTO tbl_input (
                Processor,
                Process,
                Chargeback_Date,
                Case_Json,
                Case_ID,
                Input_Identifier
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            [10, 20, "2026-05-25", '{"case": "CASE-1"}', "CASE-1", f"20_CASE-1_{status}"],
        )
        input_id = int(cursor.lastrowid)
        self.db.connection.execute(
            """
            INSERT INTO tbl_queue (Case_Details, Application_Details, Processing_Status, Bot_Comment)
            VALUES (?, ?, ?, ?)
            """,
            [input_id, 10, status, bot_comment],
        )
        self.db.connection.commit()
        return input_id

    def _insert_input(self, process: int, case_id: str, status: str | None) -> int:
        cursor = self.db.connection.execute(
            """
            INSERT INTO tbl_input (
                Processor,
                Process,
                Chargeback_Date,
                Case_Json,
                Case_ID,
                Status,
                Input_Identifier
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [10, process, "2026-05-25", "{}", case_id, status, f"{process}_{case_id}"],
        )
        self.db.connection.commit()
        return int(cursor.lastrowid)


if __name__ == "__main__":
    unittest.main()
